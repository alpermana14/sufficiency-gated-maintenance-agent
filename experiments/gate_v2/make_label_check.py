"""Gate v2 - sample sheet so an author can verify the sufficiency labels.

The 200 queries and their labels were written by the authors with the help of an AI assistant
on 16 September 2026, before any distance was computed. This sheet lets an author re-decide a
stratified sample independently, so the paper can report agreement instead of asserting that the
labels are correct. Joren et al. (ICLR 2025, Table 1, p. 5) validate their own sufficiency
autorater the same way, on 115 hand-labelled instances.

The judgement asked for is documentary, not an engineering opinion: given these twenty work
orders, does any of them contain what the question asks for?

Usage: python experiments/gate_v2/make_label_check.py [n]
Output: CAEE_R1/rating/Gate_v2_Label_Check.xlsx
"""
import json
import os
import random
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, "CAEE_R1", "rating", "Gate_v2_Label_Check.xlsx")
SEED = 20260917


def main():
    n_total = int(sys.argv[1]) if len(sys.argv) > 1 else 40
    store = json.load(open(os.path.join(HERE, "store.json"), encoding="utf-8"))
    qs = json.load(open(os.path.join(HERE, "queries.json"), encoding="utf-8"))["queries"]

    by_cat = {}
    for q in qs:
        by_cat.setdefault(q["category"], []).append(q)
    rng = random.Random(SEED)
    sample = []
    for cat, group in by_cat.items():
        take = max(2, round(n_total * len(group) / len(qs)))
        sample += rng.sample(group, min(take, len(group)))
    rng.shuffle(sample)

    wb = openpyxl.Workbook()
    note = wb.active
    note.title = "READ_ME"
    lines = [
        "Gate v2 label check (17 September 2026)",
        "",
        "Why: the 200 queries and their sufficiency labels were written by the authors with the help of an AI "
        "assistant. Nobody has checked them independently. This sheet asks one author to decide a sample again, "
        "so the paper can report agreement rather than assert correctness.",
        "",
        "What to do: read the question, read the twenty work orders in the 'store' sheet, and answer one question "
        "in column D: does any record contain the information the question asks for?",
        "  yes - a record states the specific thing asked for;",
        "  no  - the store only touches the topic without stating it, or does not cover it at all.",
        "",
        "Agreed policy (16 September 2026): a record that mentions the topic but not the detail asked for counts "
        "as NO. Two records that contradict each other still count as YES, because the information exists.",
        "",
        "This is a documentary judgement, not an engineering opinion. You are not asked whether the maintenance "
        "advice is technically right, only whether the store contains the answer.",
        "",
        "Do not open the 'hidden_reference' sheet before you finish: it holds the existing labels.",
    ]
    for i, t in enumerate(lines, start=1):
        c = note.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        if i in (1, 3, 5, 9):
            c.font = Font(bold=True)
    note.column_dimensions["A"].width = 120

    ws = wb.create_sheet("check")
    ws.append(["item", "category", "question", "your answer (yes/no)", "note (optional)"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for q in sample:
        ws.append([q["id"], q["category"], q["query"], "", ""])
    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(sample) + 1}")
    for col, w in zip("ABCDE", (10, 20, 70, 20, 40)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[3].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "C2"

    st = wb.create_sheet("store")
    st.append(["id", "date", "work order text"])
    for c in st[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in store["records"]:
        st.append([r["id"], r["date"], r["text"]])
    for col, w in zip("ABC", (10, 14, 120)):
        st.column_dimensions[col].width = w
    for row in st.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    hid = wb.create_sheet("hidden_reference")
    hid.append(["item", "existing label (sufficient)", "gold records"])
    for q in sample:
        hid.append([q["id"], "yes" if q["sufficient"] else "no", ", ".join(q["gold_records"])])
    hid.sheet_state = "hidden"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"wrote {len(sample)} items to {OUT}")
    for cat in sorted(by_cat):
        print(f"  {cat:20s} {sum(1 for q in sample if q['category'] == cat)}")


if __name__ == "__main__":
    main()
